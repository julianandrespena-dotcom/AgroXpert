
import pandas as pd
import os
import tkinter as tk
from tkinter import filedialog, messagebox


# ============================================================
# CONFIGURACIÓN PRINCIPAL
# ============================================================

# ------------------------------------------------------------
# NÚMERO DE VARIEDADES PRINCIPALES
# ------------------------------------------------------------
#
# Cambia ÚNICAMENTE este valor cuando quieras modificar
# el número de categorías principales.
#
# Ejemplos:
#
#     5  -> categorías 1, 2, 3, 4, 5 y VARIAS = 6
#     7  -> categorías 1, 2, 3, 4, 5, 6, 7 y VARIAS = 8
#    10  -> categorías 1 a 10 y VARIAS = 11
#
# ------------------------------------------------------------

NUM_CATEGORIAS_PRINCIPALES = 10


# ============================================================
# COLUMNAS QUE DEBEN CARGARSE COMO TEXTO
# ============================================================

## Historico biosal
# COLUMNAS_TEXTO = [
#     "Llave",
#     "Periodo",
#     "Archivo_origen",
#     "CODSIA",
#     "Hacienda",
#     "Tenencia",
#     "Suerte"
# ]

## crono
COLUMNAS_TEXTO = [
    "Hacienda",
    "Tenencia",
    "Suerte"
]


# ============================================================
# COLUMNAS NECESARIAS
# ============================================================

COLUMNAS_NECESARIAS = COLUMNAS_TEXTO + [
    "Variedad",
    "Area Bruta"
]


# ============================================================
# VALIDAR CONFIGURACIÓN
# ============================================================

if not isinstance(
    NUM_CATEGORIAS_PRINCIPALES,
    int
):

    raise ValueError(
        "NUM_CATEGORIAS_PRINCIPALES debe ser un número entero."
    )


if NUM_CATEGORIAS_PRINCIPALES < 1:

    raise ValueError(
        "NUM_CATEGORIAS_PRINCIPALES debe ser mayor o igual a 1."
    )


# ============================================================
# FUNCIÓN PRINCIPAL
# ============================================================

def procesar_historico():

    # --------------------------------------------------------
    # 1. SELECCIONAR ARCHIVO EXCEL
    # --------------------------------------------------------

    archivo_entrada = filedialog.askopenfilename(
        title="Seleccionar histórico de cosechas",
        filetypes=[
            ("Archivos Excel", "*.xlsx *.xls"),
            ("Excel XLSX", "*.xlsx"),
            ("Excel XLS", "*.xls"),
            ("Todos los archivos", "*.*")
        ]
    )

    if not archivo_entrada:

        print(
            "No se seleccionó ningún archivo."
        )

        return


    print("\n" + "=" * 80)
    print(
        "CLASIFICACIÓN DE VARIEDADES - "
        "HISTÓRICO DE COSECHAS"
    )
    print("=" * 80)


    print(
        "\nNúmero de categorías principales:",
        NUM_CATEGORIAS_PRINCIPALES
    )

    print(
        "Código utilizado para VARIAS:",
        NUM_CATEGORIAS_PRINCIPALES + 1
    )


    print(
        "\nArchivo seleccionado:"
    )

    print(
        archivo_entrada
    )


    # --------------------------------------------------------
    # 2. PREPARAR COLUMNAS COMO TEXTO
    # --------------------------------------------------------

    dtype_texto = {
        columna: str
        for columna in COLUMNAS_TEXTO
    }


    # --------------------------------------------------------
    # 3. LEER EXCEL
    # --------------------------------------------------------

    print(
        "\nLeyendo archivo Excel..."
    )


    try:

        df = pd.read_excel(
            archivo_entrada,
            dtype=dtype_texto
        )

    except Exception as e:

        messagebox.showerror(
            "Error",
            f"No fue posible leer el archivo Excel.\n\n{e}"
        )

        print(
            "\nERROR AL LEER EL ARCHIVO:"
        )

        print(e)

        return


    print(
        f"Registros cargados: {len(df):,}"
    )

    print(
        f"Columnas encontradas: {len(df.columns)}"
    )


    # --------------------------------------------------------
    # 4. LIMPIAR NOMBRES DE COLUMNAS
    # --------------------------------------------------------

    df.columns = (
        df.columns
        .astype(str)
        .str.strip()
    )


    # --------------------------------------------------------
    # 5. VERIFICAR COLUMNAS NECESARIAS
    # --------------------------------------------------------

    faltantes = [
        columna
        for columna in COLUMNAS_NECESARIAS
        if columna not in df.columns
    ]


    if faltantes:

        mensaje = (
            "El archivo no contiene las siguientes "
            "columnas necesarias:\n\n"
            +
            "\n".join(
                f"• {columna}"
                for columna in faltantes
            )
        )


        print(
            "\n" + mensaje
        )


        messagebox.showerror(
            "Columnas faltantes",
            mensaje
        )

        return


    # --------------------------------------------------------
    # 6. ASEGURAR COLUMNAS DE CÓDIGOS COMO TEXTO
    # --------------------------------------------------------

    print(
        "\nConfigurando columnas de códigos como texto..."
    )


    for columna in COLUMNAS_TEXTO:

        df[columna] = (
            df[columna]
            .astype("string")
            .str.strip()
        )


    # --------------------------------------------------------
    # 7. LIMPIAR VARIEDAD
    # --------------------------------------------------------

    df["Variedad"] = (
        df["Variedad"]
        .astype("string")
        .str.strip()
    )


    # --------------------------------------------------------
    # 8. CONVERTIR AREA BRUTA A NUMÉRICO
    # --------------------------------------------------------

    print(
        "Procesando Area Bruta..."
    )


    df["Area Bruta"] = pd.to_numeric(
        df["Area Bruta"],
        errors="coerce"
    )


    registros_area_invalida = (
        df["Area Bruta"]
        .isna()
        .sum()
    )


    if registros_area_invalida > 0:

        print(
            f"\nADVERTENCIA:"
            f"\n{registros_area_invalida:,} registros "
            "tienen Area Bruta vacía o no numérica."
        )


    # --------------------------------------------------------
    # 9. CREAR DATAFRAME PARA ANALIZAR VARIEDADES
    # --------------------------------------------------------

    df_variedades = df[
        df["Variedad"].notna()
        &
        (df["Variedad"].str.strip() != "")
        &
        df["Area Bruta"].notna()
    ].copy()


    if df_variedades.empty:

        mensaje = (
            "No existen registros válidos para realizar "
            "el análisis de variedades."
        )


        messagebox.showerror(
            "Sin datos",
            mensaje
        )

        return


    # --------------------------------------------------------
    # 10. SUMAR AREA BRUTA POR VARIEDAD
    # --------------------------------------------------------

    print(
        "\nCalculando Area Bruta acumulada por variedad..."
    )


    resumen_variedades = (
        df_variedades
        .groupby(
            "Variedad",
            as_index=False
        )["Area Bruta"]
        .sum()
        .rename(
            columns={
                "Area Bruta":
                "Area_Bruta_Total"
            }
        )
    )


    # --------------------------------------------------------
    # 11. ORDENAR DE MAYOR A MENOR
    # --------------------------------------------------------

    resumen_variedades = (
        resumen_variedades
        .sort_values(
            by="Area_Bruta_Total",
            ascending=False
        )
        .reset_index(
            drop=True
        )
    )


    # --------------------------------------------------------
    # 12. ASIGNAR RANKING
    # --------------------------------------------------------

    resumen_variedades["Ranking"] = (
        resumen_variedades.index + 1
    )


    # --------------------------------------------------------
    # 13. VERIFICAR QUE EXISTAN SUFICIENTES VARIEDADES
    # --------------------------------------------------------

    total_variedades = (
        len(resumen_variedades)
    )


    cantidad_principales = min(
        NUM_CATEGORIAS_PRINCIPALES,
        total_variedades
    )


    # --------------------------------------------------------
    # 14. IDENTIFICAR VARIEDADES PRINCIPALES
    # --------------------------------------------------------

    top_variedades = (
        resumen_variedades
        .head(cantidad_principales)
        ["Variedad"]
        .tolist()
    )


    # --------------------------------------------------------
    # 15. MOSTRAR VARIEDADES PRINCIPALES
    # --------------------------------------------------------

    print(
        "\n" + "=" * 80
    )

    print(
        f"{cantidad_principales} VARIEDADES PRINCIPALES"
    )

    print(
        "=" * 80
    )


    for codigo, variedad in enumerate(
        top_variedades,
        start=1
    ):

        fila = resumen_variedades[
            resumen_variedades["Variedad"]
            == variedad
        ]


        area = fila[
            "Area_Bruta_Total"
        ].iloc[0]


        print(
            f"{codigo}. "
            f"{variedad} "
            f"-> {area:,.2f} ha"
        )


    # --------------------------------------------------------
    # 16. CÓDIGO PARA VARIAS
    # --------------------------------------------------------

    codigo_varias = (
        NUM_CATEGORIAS_PRINCIPALES + 1
    )


    # --------------------------------------------------------
    # 17. CREAR DICCIONARIO DE CÓDIGOS
    # --------------------------------------------------------

    diccionario_codigo = {
        variedad: codigo
        for codigo, variedad
        in enumerate(
            top_variedades,
            start=1
        )
    }


    # --------------------------------------------------------
    # 18. CREAR Variedad_Codigo
    # --------------------------------------------------------

    print(
        "\nAsignando códigos de clasificación..."
    )


    df["Variedad_Codigo"] = (
        df["Variedad"]
        .map(diccionario_codigo)
        .fillna(codigo_varias)
        .astype("Int64")
    )


    # --------------------------------------------------------
    # 19. CREAR Variedad_Clasificacion
    # --------------------------------------------------------

    def clasificar_variedad(variedad):

        if pd.isna(variedad):

            return "SIN VARIEDAD"


        variedad = str(
            variedad
        ).strip()


        if variedad == "":

            return "SIN VARIEDAD"


        if variedad in diccionario_codigo:

            return variedad


        return "VARIAS"


    df["Variedad_Clasificacion"] = (
        df["Variedad"]
        .apply(
            clasificar_variedad
        )
    )


    # --------------------------------------------------------
    # 20. CREAR RESUMEN DE CLASIFICACIÓN
    # --------------------------------------------------------

    distribucion = (
        df
        .groupby(
            [
                "Variedad_Codigo",
                "Variedad_Clasificacion"
            ],
            dropna=False
        )
        .agg(
            Registros=(
                "Variedad",
                "size"
            ),

            Area_Bruta=(
                "Area Bruta",
                "sum"
            )
        )
        .reset_index()
        .sort_values(
            by="Variedad_Codigo"
        )
    )


    # --------------------------------------------------------
    # 21. AGREGAR CLASIFICACIÓN AL RESUMEN
    # --------------------------------------------------------

    resumen_variedades[
        "Codigo_Clasificacion"
    ] = (
        resumen_variedades["Variedad"]
        .map(diccionario_codigo)
        .fillna(codigo_varias)
        .astype(int)
    )


    resumen_variedades[
        "Clasificacion"
    ] = (
        resumen_variedades["Variedad"]
        .apply(
            lambda variedad:
                variedad
                if variedad in diccionario_codigo
                else "VARIAS"
        )
    )


    # --------------------------------------------------------
    # 22. REORDENAR RESUMEN
    # --------------------------------------------------------

    resumen_variedades = (
        resumen_variedades[
            [
                "Ranking",
                "Variedad",
                "Area_Bruta_Total",
                "Codigo_Clasificacion",
                "Clasificacion"
            ]
        ]
    )


    # --------------------------------------------------------
    # 23. REORDENAR COLUMNAS DEL HISTÓRICO
    # --------------------------------------------------------

    columnas_originales = [
        columna
        for columna in df.columns
        if columna not in [
            "Variedad_Codigo",
            "Variedad_Clasificacion"
        ]
    ]


    posicion_variedad = (
        columnas_originales
        .index("Variedad")
    )


    columnas_finales = (
        columnas_originales[
            :posicion_variedad + 1
        ]
        +
        [
            "Variedad_Codigo",
            "Variedad_Clasificacion"
        ]
        +
        columnas_originales[
            posicion_variedad + 1:
        ]
    )


    df = df[
        columnas_finales
    ]


    # --------------------------------------------------------
    # 24. DEFINIR ARCHIVO DE SALIDA
    # --------------------------------------------------------

    carpeta = os.path.dirname(
        archivo_entrada
    )


    nombre_archivo = os.path.basename(
        archivo_entrada
    )


    nombre_sin_extension = (
        os.path.splitext(
            nombre_archivo
        )[0]
    )


    archivo_salida = os.path.join(
        carpeta,
        f"{nombre_sin_extension}_CLASIFICADO.xlsx"
    )


    # --------------------------------------------------------
    # 25. EVITAR SOBRESCRIBIR ARCHIVOS EXISTENTES
    # --------------------------------------------------------

    if os.path.exists(
        archivo_salida
    ):

        contador = 1


        while True:

            archivo_alternativo = os.path.join(
                carpeta,
                f"{nombre_sin_extension}_CLASIFICADO_{contador}.xlsx"
            )


            if not os.path.exists(
                archivo_alternativo
            ):

                archivo_salida = (
                    archivo_alternativo
                )

                break


            contador += 1


    # --------------------------------------------------------
    # 26. GUARDAR EXCEL
    # --------------------------------------------------------

    print(
        "\nGenerando archivo de salida..."
    )


    try:

        with pd.ExcelWriter(
            archivo_salida,
            engine="openpyxl"
        ) as writer:

            # --------------------------------------------
            # HOJA 1
            # --------------------------------------------

            df.to_excel(
                writer,
                sheet_name="Historico_Clasificado",
                index=False
            )


            # --------------------------------------------
            # HOJA 2
            # --------------------------------------------

            resumen_variedades.to_excel(
                writer,
                sheet_name="Resumen_Variedades",
                index=False
            )


            # --------------------------------------------
            # HOJA 3
            # --------------------------------------------

            distribucion.to_excel(
                writer,
                sheet_name="Resumen_Clasificacion",
                index=False
            )


    except Exception as e:

        print(
            "\nERROR AL GUARDAR:"
        )

        print(e)


        messagebox.showerror(
            "Error al guardar",
            f"No fue posible guardar el archivo.\n\n{e}"
        )

        return


    # --------------------------------------------------------
    # 27. FORMATEAR EXCEL
    # --------------------------------------------------------

    try:

        from openpyxl import load_workbook
        from openpyxl.styles import Font, Alignment
        from openpyxl.utils import get_column_letter


        wb = load_workbook(
            archivo_salida
        )


        # --------------------------------------------
        # FORMATO GENERAL
        # --------------------------------------------

        for ws in wb.worksheets:

            # Congelar encabezado
            ws.freeze_panes = "A2"


            # Autofiltro
            if ws.max_row > 1:

                ws.auto_filter.ref = (
                    ws.dimensions
                )


            # Encabezados
            for cell in ws[1]:

                cell.font = Font(
                    bold=True
                )

                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center"
                )


            # Ajustar ancho de columnas
            for numero_columna in range(
                1,
                ws.max_column + 1
            ):

                letra = get_column_letter(
                    numero_columna
                )


                max_length = 0


                for celda in ws[letra]:

                    try:

                        longitud = len(
                            str(
                                celda.value
                            )
                        )


                        if longitud > max_length:

                            max_length = (
                                longitud
                            )

                    except:

                        pass


                ancho = min(
                    max(
                        max_length + 2,
                        10
                    ),
                    40
                )


                ws.column_dimensions[
                    letra
                ].width = ancho


        # --------------------------------------------
        # FORMATO AREA BRUTA
        # --------------------------------------------

        ws = wb[
            "Historico_Clasificado"
        ]


        columna_area = None


        for cell in ws[1]:

            if cell.value == "Area Bruta":

                columna_area = (
                    cell.column
                )

                break


        if columna_area is not None:

            for fila in range(
                2,
                ws.max_row + 1
            ):

                ws.cell(
                    fila,
                    columna_area
                ).number_format = (
                    '#,##0.00'
                )


        # --------------------------------------------
        # FORMATO RESUMEN VARIEDADES
        # --------------------------------------------

        ws = wb[
            "Resumen_Variedades"
        ]


        columna_resumen = None


        for cell in ws[1]:

            if cell.value == "Area_Bruta_Total":

                columna_resumen = (
                    cell.column
                )

                break


        if columna_resumen is not None:

            for fila in range(
                2,
                ws.max_row + 1
            ):

                ws.cell(
                    fila,
                    columna_resumen
                ).number_format = (
                    '#,##0.00'
                )


        # --------------------------------------------
        # FORMATO RESUMEN CLASIFICACIÓN
        # --------------------------------------------

        ws = wb[
            "Resumen_Clasificacion"
        ]


        columna_distribucion = None


        for cell in ws[1]:

            if cell.value == "Area_Bruta":

                columna_distribucion = (
                    cell.column
                )

                break


        if columna_distribucion is not None:

            for fila in range(
                2,
                ws.max_row + 1
            ):

                ws.cell(
                    fila,
                    columna_distribucion
                ).number_format = (
                    '#,##0.00'
                )


        # Guardar
        wb.save(
            archivo_salida
        )


    except Exception as e:

        print(
            "\nADVERTENCIA:"
        )

        print(
            "El archivo fue creado, "
            "pero no se pudo aplicar "
            "todo el formato."
        )

        print(e)


    # --------------------------------------------------------
    # 28. RESULTADO FINAL
    # --------------------------------------------------------

    print(
        "\n" + "=" * 80
    )

    print(
        "PROCESO TERMINADO CORRECTAMENTE"
    )

    print(
        "=" * 80
    )


    print(
        f"\nRegistros procesados: "
        f"{len(df):,}"
    )


    print(
        f"Variedades encontradas: "
        f"{total_variedades:,}"
    )


    print(
        f"Categorías principales: "
        f"{cantidad_principales}"
    )


    print(
        f"Código VARIAS: "
        f"{codigo_varias}"
    )


    print(
        "\nClasificación utilizada:"
    )


    for codigo, variedad in enumerate(
        top_variedades,
        start=1
    ):

        print(
            f"  {codigo} = {variedad}"
        )


    print(
        f"  {codigo_varias} = VARIAS"
    )


    print(
        "\nArchivo generado:"
    )


    print(
        archivo_salida
    )


    print(
        "\nHojas generadas:"
    )


    print(
        "  1. Historico_Clasificado"
    )


    print(
        "  2. Resumen_Variedades"
    )


    print(
        "  3. Resumen_Clasificacion"
    )


    print(
        "\n" + "=" * 80
    )


    # --------------------------------------------------------
    # 29. MENSAJE FINAL
    # --------------------------------------------------------

    messagebox.showinfo(
        "Proceso terminado",
        "El histórico fue procesado correctamente.\n\n"
        f"Registros: {len(df):,}\n"
        f"Variedades encontradas: {total_variedades:,}\n\n"
        f"Categorías principales: "
        f"{cantidad_principales}\n"
        f"VARIAS = código {codigo_varias}\n\n"
        f"Archivo generado:\n{archivo_salida}"
    )


# ============================================================
# INICIAR PROGRAMA
# ============================================================

if __name__ == "__main__":

    root = tk.Tk()

    # Ocultar ventana principal de Tkinter
    root.withdraw()

    # Ejecutar proceso
    procesar_historico()

    # Cerrar Tkinter
    root.destroy()



